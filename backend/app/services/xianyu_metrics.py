from __future__ import annotations

import hashlib
import json
from collections import defaultdict
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Account, User, XianyuProduct, XianyuProductDailyMetric
from app.services.audit import write_audit

REPORT_SHEET = "商品明细表"
MAX_REPORT_BYTES = 5 * 1024 * 1024
MAX_REPORT_ROWS = 5_000
EXPECTED_HEADERS = (
    "商品 ID",
    "商品名称",
    "商品照片链接",
    "渠道行业类目名称",
    "渠道一级类目名称",
    "渠道二级类目名称",
    "渠道三级类目名称",
    "渠道四级类目名称",
    "商品价格",
    "累计上架天数",
    "上架日期",
    "商品曝光次数",
    "商品曝光人数",
    "商品浏览次数",
    "商品浏览人数",
    "询单人数",
    "支付人数",
    "支付订单数",
    "支付金额",
    "点击支付转化率",
    "发起退款人数",
    "发起退款订单数",
    "发起退款金额",
    "成功退款人数",
    "成功退款订单数",
    "成功退款金额",
)
COUNT_HEADERS = {
    "累计上架天数",
    "商品曝光次数",
    "商品曝光人数",
    "商品浏览次数",
    "商品浏览人数",
    "询单人数",
    "支付人数",
    "支付订单数",
    "发起退款人数",
    "发起退款订单数",
    "成功退款人数",
    "成功退款订单数",
}
MONEY_HEADERS = {"商品价格", "支付金额", "发起退款金额", "成功退款金额"}


def _account_for_user(db: Session, user: User) -> Account:
    account = db.scalar(
        select(Account)
        .where(Account.platform == "XIANYU", Account.owner_user_id == user.id)
        .order_by(Account.id)
    )
    if account is None:
        raise LookupError("请先导入闲鱼商品快照，建立账号归属后再导入经营报表")
    return account


def _text(value, *, field: str, maximum: int) -> str:
    normalized = str(value or "").strip()
    if not normalized:
        raise ValueError(f"{field} 不能为空")
    if len(normalized) > maximum or any(ord(character) < 32 for character in normalized):
        raise ValueError(f"{field} 格式无效")
    return normalized


def _product_id(value) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    normalized = str(value or "").strip()
    if not normalized.isdigit() or not 6 <= len(normalized) <= 20:
        raise ValueError("商品 ID 必须为 6–20 位数字")
    return normalized


def _decimal(value, *, field: str, optional: bool = False) -> Decimal | None:
    if value in {None, "", "-", "--"}:
        return None if optional else Decimal("0")
    normalized = str(value).strip().replace(",", "").replace("¥", "").replace("￥", "")
    is_percent = normalized.endswith("%")
    if is_percent:
        normalized = normalized[:-1]
    try:
        number = Decimal(normalized)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{field} 不是有效数字") from exc
    if is_percent:
        number /= Decimal("100")
    if number < 0:
        raise ValueError(f"{field} 不能为负数")
    return number


def _count(value, *, field: str) -> int:
    number = _decimal(value, field=field)
    if number != number.to_integral_value():
        raise ValueError(f"{field} 必须为整数")
    return int(number)


def _listed_at(value) -> datetime | None:
    if value in {None, "", "-", "--"}:
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=value.tzinfo or UTC)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=UTC)
    try:
        parsed = datetime.fromisoformat(str(value).strip().replace("/", "-"))
    except ValueError as exc:
        raise ValueError("上架日期格式无效") from exc
    return parsed.replace(tzinfo=parsed.tzinfo or UTC)


def parse_seller_workbench_xlsx(content: bytes) -> list[dict]:
    if not content:
        raise ValueError("经营报表为空")
    if len(content) > MAX_REPORT_BYTES:
        raise ValueError("经营报表超过 5 MB 限制")
    if not content.startswith(b"PK"):
        raise ValueError("只接受闲鱼经营罗盘导出的 XLSX 文件")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError("无法读取 XLSX，请重新从闲鱼经营罗盘下载") from exc
    try:
        if REPORT_SHEET not in workbook.sheetnames:
            raise ValueError(f"报表缺少工作表：{REPORT_SHEET}")
        sheet = workbook[REPORT_SHEET]
        if sheet.max_column != len(EXPECTED_HEADERS) or sheet.max_row > MAX_REPORT_ROWS + 1:
            raise ValueError("报表列数不符或商品行数超过 5000")
        header_cells = next(sheet.iter_rows(min_row=1, max_row=1))
        headers = tuple(str(cell.value or "").strip() for cell in header_cells)
        if headers != EXPECTED_HEADERS:
            raise ValueError("报表字段与闲鱼经营罗盘商品明细模板不一致")

        records: list[dict] = []
        seen_ids: set[str] = set()
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
            if all(cell.value in {None, ""} for cell in cells):
                continue
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError(f"第 {row_number} 行包含公式，拒绝导入可执行单元格")
            raw = dict(zip(headers, (cell.value for cell in cells), strict=True))
            try:
                external_product_id = _product_id(raw["商品 ID"])
                if external_product_id in seen_ids:
                    raise ValueError("报表内商品 ID 重复")
                seen_ids.add(external_product_id)
                record = {
                    "external_product_id": external_product_id,
                    "listing_title": _text(raw["商品名称"], field="商品名称", maximum=500),
                    "listing_price": _decimal(raw["商品价格"], field="商品价格"),
                    "listing_days": _count(raw["累计上架天数"], field="累计上架天数"),
                    "listed_at": _listed_at(raw["上架日期"]),
                    "category_path": [
                        str(raw[field]).strip()
                        for field in EXPECTED_HEADERS[3:8]
                        if raw[field] not in {None, "", "-", "--"}
                    ],
                    "exposures": _count(raw["商品曝光次数"], field="商品曝光次数"),
                    "exposed_users": _count(raw["商品曝光人数"], field="商品曝光人数"),
                    "views": _count(raw["商品浏览次数"], field="商品浏览次数"),
                    "viewers": _count(raw["商品浏览人数"], field="商品浏览人数"),
                    "inquiries": _count(raw["询单人数"], field="询单人数"),
                    "paid_users": _count(raw["支付人数"], field="支付人数"),
                    "paid_orders": _count(raw["支付订单数"], field="支付订单数"),
                    "paid_amount": _decimal(raw["支付金额"], field="支付金额"),
                    "browse_pay_conversion": _decimal(
                        raw["点击支付转化率"], field="点击支付转化率", optional=True
                    ),
                    "refund_requested_users": _count(
                        raw["发起退款人数"], field="发起退款人数"
                    ),
                    "refund_requested_orders": _count(
                        raw["发起退款订单数"], field="发起退款订单数"
                    ),
                    "refund_requested_amount": _decimal(
                        raw["发起退款金额"], field="发起退款金额"
                    ),
                    "refund_success_users": _count(
                        raw["成功退款人数"], field="成功退款人数"
                    ),
                    "refund_success_orders": _count(
                        raw["成功退款订单数"], field="成功退款订单数"
                    ),
                    "refund_success_amount": _decimal(
                        raw["成功退款金额"], field="成功退款金额"
                    ),
                }
            except ValueError as exc:
                raise ValueError(f"第 {row_number} 行：{exc}") from exc
            if record["views"] > record["exposures"] and record["exposures"] > 0:
                raise ValueError(f"第 {row_number} 行：浏览次数不能大于曝光次数")
            records.append(record)
        if not records:
            raise ValueError("报表没有商品数据")
        return records
    finally:
        workbook.close()


def import_seller_workbench_xlsx(
    db: Session,
    *,
    content: bytes,
    metric_date: date,
    user: User,
    correlation_id: str,
) -> dict:
    today = datetime.now(ZoneInfo("Asia/Shanghai")).date()
    if metric_date >= today:
        raise ValueError("经营罗盘为 T+1 数据，统计日期必须早于今天")
    if metric_date < today - timedelta(days=365):
        raise ValueError("统计日期超过一年，请检查报表周期")
    records = parse_seller_workbench_xlsx(content)
    account = _account_for_user(db, user)
    source_hash = hashlib.sha256(content).hexdigest()
    products = {
        item.external_product_id: item
        for item in db.scalars(
            select(XianyuProduct).where(XianyuProduct.account_id == account.id)
        ).all()
        if item.external_product_id
    }
    managed_external_ids = {
        external_id
        for external_id, item in products.items()
        if item.product_id is not None
    }
    existing = {
        item.external_product_id: item
        for item in db.scalars(
            select(XianyuProductDailyMetric).where(
                XianyuProductDailyMetric.account_id == account.id,
                XianyuProductDailyMetric.metric_date == metric_date,
            )
        ).all()
    }
    created = 0
    updated = 0
    for record in records:
        external_product_id = record["external_product_id"]
        row = existing.get(external_product_id)
        if row is None:
            row = XianyuProductDailyMetric(
                account_id=account.id,
                external_product_id=external_product_id,
                metric_date=metric_date,
                source_hash=source_hash,
                imported_by_user_id=user.id,
            )
            db.add(row)
            created += 1
        else:
            updated += 1
        row.xianyu_product_id = getattr(products.get(external_product_id), "id", None)
        row.source = "seller-workbench-xlsx"
        row.source_hash = source_hash
        row.imported_by_user_id = user.id
        for key, value in record.items():
            if key != "external_product_id":
                setattr(row, key, value)
        row.raw_snapshot = {
            "metric_date": metric_date.isoformat(),
            "source": "seller-workbench-xlsx",
            "category_path": record["category_path"],
        }
    write_audit(
        db,
        action="XIANYU_DAILY_METRICS_IMPORTED",
        entity_type="ACCOUNT",
        entity_id=account.id,
        actor_user_id=user.id,
        actor_type="USER",
        after_data={
            "metric_date": metric_date.isoformat(),
            "rows": len(records),
            "created": created,
            "updated": updated,
            "managed_products": sum(
                record["external_product_id"] in managed_external_ids for record in records
            ),
            "source_hash": source_hash,
        },
        correlation_id=correlation_id,
    )
    db.commit()
    return {
        "account_id": account.id,
        "metric_date": metric_date.isoformat(),
        "rows": len(records),
        "created": created,
        "updated": updated,
        "managed_products": sum(
            record["external_product_id"] in managed_external_ids for record in records
        ),
        "source_hash": source_hash,
    }


def ingest_browser_metrics(
    db: Session,
    *,
    records: list[dict],
    metric_date: date,
    user: User,
    correlation_id: str,
) -> dict:
    today = datetime.now(ZoneInfo("Asia/Shanghai")).date()
    if metric_date >= today or metric_date < today - timedelta(days=7):
        raise ValueError("浏览器日报日期必须是最近 7 天内的 T+1 日期")
    if not 1 <= len(records) <= 500:
        raise ValueError("浏览器日报商品数必须在 1–500 之间")
    account = _account_for_user(db, user)
    source_hash = hashlib.sha256(
        json.dumps(records, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")
    ).hexdigest()
    products = {
        item.external_product_id: item
        for item in db.scalars(
            select(XianyuProduct).where(XianyuProduct.account_id == account.id)
        ).all()
        if item.external_product_id
    }
    existing = {
        item.external_product_id: item
        for item in db.scalars(
            select(XianyuProductDailyMetric).where(
                XianyuProductDailyMetric.account_id == account.id,
                XianyuProductDailyMetric.metric_date == metric_date,
            )
        ).all()
    }
    created = 0
    updated = 0
    seen: set[str] = set()
    for record in records:
        external_product_id = _product_id(record.get("external_product_id"))
        if external_product_id in seen:
            raise ValueError("浏览器日报包含重复商品 ID")
        seen.add(external_product_id)
        product = products.get(external_product_id)
        row = existing.get(external_product_id)
        if row is None:
            row = XianyuProductDailyMetric(
                account_id=account.id,
                external_product_id=external_product_id,
                metric_date=metric_date,
                source_hash=source_hash,
                imported_by_user_id=user.id,
            )
            db.add(row)
            created += 1
        else:
            updated += 1
        row.xianyu_product_id = product.id if product else None
        row.listing_title = _text(
            record.get("listing_title"), field="商品名称", maximum=500
        )
        row.listing_price = _decimal(record.get("listing_price"), field="商品价格")
        if product and product.created_at:
            created_date = _utc_date(product.created_at)
            row.listing_days = max(0, (metric_date - created_date).days)
        else:
            row.listing_days = 0
        row.listed_at = product.created_at if product else None
        row.category_path = []
        for field in (
            "exposures",
            "exposed_users",
            "views",
            "viewers",
            "inquiries",
            "paid_users",
            "paid_orders",
            "refund_requested_users",
            "refund_requested_orders",
            "refund_success_users",
            "refund_success_orders",
        ):
            setattr(row, field, _count(record.get(field), field=field))
        for field in (
            "paid_amount",
            "refund_requested_amount",
            "refund_success_amount",
        ):
            setattr(row, field, _decimal(record.get(field), field=field))
        row.browse_pay_conversion = _decimal(
            record.get("browse_pay_conversion"),
            field="浏览支付转化率",
            optional=True,
        )
        if row.views > row.exposures and row.exposures > 0:
            raise ValueError(f"商品 {external_product_id} 的浏览次数大于曝光次数")
        row.source = "seller-workbench-browser-bridge"
        row.source_hash = source_hash
        row.imported_by_user_id = user.id
        row.raw_snapshot = {
            "metric_date": metric_date.isoformat(),
            "source": "seller-workbench-browser-bridge",
        }
    managed_ids = {
        external_id
        for external_id, product in products.items()
        if product.product_id is not None
    }
    write_audit(
        db,
        action="XIANYU_DAILY_METRICS_CAPTURED",
        entity_type="ACCOUNT",
        entity_id=account.id,
        actor_type="BROWSER_BRIDGE",
        after_data={
            "metric_date": metric_date.isoformat(),
            "rows": len(records),
            "created": created,
            "updated": updated,
            "managed_products": len(seen & managed_ids),
            "source_hash": source_hash,
        },
        correlation_id=correlation_id,
    )
    return {
        "account_id": account.id,
        "metric_date": metric_date.isoformat(),
        "rows": len(records),
        "created": created,
        "updated": updated,
        "managed_products": len(seen & managed_ids),
        "source_hash": source_hash,
    }


def _utc_date(value: datetime) -> date:
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(UTC).date()


def _ratio(numerator: int, denominator: int) -> Decimal:
    if denominator <= 0:
        return Decimal("0")
    return (Decimal(numerator) / Decimal(denominator)).quantize(Decimal("0.0001"))


def diagnose_product(*, data_days: int, listing_days: int, metrics: dict) -> dict:
    exposures = metrics["exposures"]
    views = metrics["views"]
    inquiries = metrics["inquiries"]
    paid_orders = metrics["paid_orders"]
    refund_orders = metrics["refund_success_orders"]
    ctr = _ratio(views, exposures)
    inquiry_rate = _ratio(inquiries, views)
    pay_rate = _ratio(paid_orders, views)
    refund_rate = _ratio(refund_orders, paid_orders)

    if paid_orders >= 3 and refund_rate >= Decimal("0.10"):
        return {
            "code": "HIGH_REFUND",
            "label": "退款风险",
            "severity": "critical",
            "actions": [
                "暂停放量与自动调价",
                "核对规格、质量、包装和描述一致性",
                "人工复盘退款原因",
            ],
        }
    if listing_days >= 3 and exposures < 30:
        return {
            "code": "NO_EXPOSURE",
            "label": "曝光不足",
            "severity": "high",
            "actions": ["核对类目是否准确", "补充搜索关键词但避免堆砌", "检查账号垂直度与发布时间"],
        }
    if data_days < 3 or exposures < 60:
        signal = "当前点击偏低" if exposures >= 30 and ctr < Decimal("0.02") else "数据量不足"
        return {
            "code": "COLLECTING_DATA",
            "label": f"采集中 · {signal}",
            "severity": "info",
            "actions": [
                "保持价格不变",
                "连续采集满 3 个完整日",
                "达到 60 次曝光后再决定是否改主图或标题",
            ],
        }
    if ctr < Decimal("0.02"):
        return {
            "code": "LOW_CLICK",
            "label": "点击率偏低",
            "severity": "high",
            "actions": [
                "优先更换首图并保留原价",
                "标题前 20 字突出品类和核心规格",
                "每次只改一个变量并观察 72 小时",
            ],
        }
    if views >= 20 and inquiry_rate < Decimal("0.03"):
        return {
            "code": "LOW_INTENT",
            "label": "浏览后意向偏低",
            "severity": "medium",
            "actions": [
                "补齐尺寸、材质和使用场景",
                "检查价格与同类差距",
                "强化现货、包装和售后事实",
            ],
        }
    if inquiries >= 3 and paid_orders == 0:
        return {
            "code": "INQUIRY_NOT_PAID",
            "label": "询单未成交",
            "severity": "medium",
            "actions": ["缩短首响时间", "统一回答库存、运费和发货时效", "核对议价底线与常见异议"],
        }
    if ctr >= Decimal("0.05") or pay_rate > 0:
        return {
            "code": "HEALTHY",
            "label": "表现健康",
            "severity": "success",
            "actions": ["保持当前首图与标题", "继续监控成交和退款", "优先复制有效素材结构"],
        }
    return {
        "code": "OBSERVE",
        "label": "继续观察",
        "severity": "info",
        "actions": ["保持当前变量", "继续采集至 7 天", "优先优化回复与详情完整度"],
    }


def metrics_overview(db: Session, *, user: User, days: int = 7) -> dict:
    if not 1 <= days <= 90:
        raise ValueError("统计周期必须在 1–90 天之间")
    account = _account_for_user(db, user)
    managed_products = db.scalars(
        select(XianyuProduct).where(
            XianyuProduct.account_id == account.id,
            XianyuProduct.product_id.is_not(None),
            XianyuProduct.status.in_(["ACTIVE", "PAUSED"]),
        )
    ).all()
    managed_ids = {item.id for item in managed_products}
    latest_date = db.scalar(
        select(XianyuProductDailyMetric.metric_date)
        .where(XianyuProductDailyMetric.account_id == account.id)
        .order_by(XianyuProductDailyMetric.metric_date.desc())
        .limit(1)
    )
    if latest_date is None:
        return {
            "latest_date": None,
            "period_start": None,
            "period_end": None,
            "data_days": 0,
            "summary": {
                "products": 0,
                "exposures": 0,
                "views": 0,
                "ctr": "0.0000",
                "inquiries": 0,
                "paid_orders": 0,
                "paid_amount": "0.00",
            },
            "daily": [],
            "products": [],
            "strategy": {
                "observation_days": 3,
                "minimum_exposures": 60,
                "price_change_days": 7,
                "note": "暂无官方经营数据，请导入经营罗盘近1天报表。",
            },
        }

    period_start = latest_date - timedelta(days=days - 1)
    rows = db.scalars(
        select(XianyuProductDailyMetric)
        .where(
            XianyuProductDailyMetric.account_id == account.id,
            XianyuProductDailyMetric.metric_date >= period_start,
            XianyuProductDailyMetric.metric_date <= latest_date,
        )
        .order_by(XianyuProductDailyMetric.metric_date, XianyuProductDailyMetric.id)
    ).all()
    daily_groups: dict[date, list[XianyuProductDailyMetric]] = defaultdict(list)
    product_groups: dict[str, list[XianyuProductDailyMetric]] = defaultdict(list)
    for row in rows:
        daily_groups[row.metric_date].append(row)
        product_groups[row.external_product_id].append(row)

    sum_fields = (
        "exposures",
        "exposed_users",
        "views",
        "viewers",
        "inquiries",
        "paid_users",
        "paid_orders",
        "refund_requested_orders",
        "refund_success_orders",
    )

    def aggregate(items):
        result = {field: sum(getattr(item, field) for item in items) for field in sum_fields}
        result["paid_amount"] = sum((item.paid_amount for item in items), Decimal("0"))
        result["ctr"] = _ratio(result["views"], result["exposures"])
        result["inquiry_rate"] = _ratio(result["inquiries"], result["views"])
        result["pay_rate"] = _ratio(result["paid_orders"], result["views"])
        return result

    total = aggregate(rows)
    managed_total = aggregate([row for row in rows if row.xianyu_product_id in managed_ids])
    daily = []
    for metric_day, items in daily_groups.items():
        metrics = aggregate(items)
        daily.append(
            {
                "date": metric_day.isoformat(),
                "products": len(items),
                "exposures": metrics["exposures"],
                "views": metrics["views"],
                "ctr": str(metrics["ctr"]),
                "inquiries": metrics["inquiries"],
                "paid_orders": metrics["paid_orders"],
            }
        )

    products = []
    for external_product_id, items in product_groups.items():
        newest = max(items, key=lambda item: item.metric_date)
        metrics = aggregate(items)
        managed = newest.xianyu_product_id in managed_ids
        diagnosis = diagnose_product(
            data_days=len({item.metric_date for item in items}),
            listing_days=newest.listing_days,
            metrics=metrics,
        )
        products.append(
            {
                "external_product_id": external_product_id,
                "xianyu_product_id": newest.xianyu_product_id,
                "managed": managed,
                "title": newest.listing_title,
                "price": str(newest.listing_price),
                "listing_days": newest.listing_days,
                "data_days": len({item.metric_date for item in items}),
                "exposures": metrics["exposures"],
                "views": metrics["views"],
                "ctr": str(metrics["ctr"]),
                "inquiries": metrics["inquiries"],
                "inquiry_rate": str(metrics["inquiry_rate"]),
                "paid_orders": metrics["paid_orders"],
                "paid_amount": str(metrics["paid_amount"]),
                "diagnosis": diagnosis,
            }
        )
    reported_external_ids = set(product_groups)
    for item in managed_products:
        if not item.external_product_id or item.external_product_id in reported_external_ids:
            continue
        products.append(
            {
                "external_product_id": item.external_product_id,
                "xianyu_product_id": item.id,
                "managed": True,
                "title": item.published_title or "未命名商品",
                "price": str(item.published_price or Decimal("0")),
                "listing_days": 0,
                "data_days": 0,
                "exposures": 0,
                "views": 0,
                "ctr": "0.0000",
                "inquiries": 0,
                "inquiry_rate": "0.0000",
                "paid_orders": 0,
                "paid_amount": "0.00",
                "diagnosis": {
                    "code": "AWAITING_T1",
                    "label": "等待 T+1",
                    "severity": "info",
                    "actions": [
                        "次日导入近1天报表",
                        "保持首图、标题和价格不变",
                        "采满 3 个完整日后首次评估",
                    ],
                },
            }
        )
    products.sort(key=lambda item: (not item["managed"], item["ctr"], -item["exposures"]))
    return {
        "latest_date": latest_date.isoformat(),
        "period_start": period_start.isoformat(),
        "period_end": latest_date.isoformat(),
        "data_days": len(daily_groups),
        "summary": {
            "products": len(product_groups),
            "managed_products": sum(item["managed"] for item in products),
            "awaiting_t1": sum(
                item["diagnosis"]["code"] == "AWAITING_T1" for item in products
            ),
            "exposures": total["exposures"],
            "views": total["views"],
            "ctr": str(total["ctr"]),
            "managed_exposures": managed_total["exposures"],
            "managed_views": managed_total["views"],
            "managed_ctr": str(managed_total["ctr"]),
            "inquiries": total["inquiries"],
            "paid_orders": total["paid_orders"],
            "paid_amount": str(total["paid_amount"]),
        },
        "daily": daily,
        "products": products,
        "strategy": {
            "observation_days": 3,
            "minimum_exposures": 60,
            "price_change_days": 7,
            "note": (
                "主图或标题每次只改一个变量，观察 72 小时；"
                "价格需至少 7 天证据且不得低于最低售价。"
            ),
        },
    }
